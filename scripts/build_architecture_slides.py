#!/usr/bin/env python3
"""
Build the DengueCAD architecture slide deck.

Reads the rendered diagrams in docs/architecture/*.png (run
render_architecture_diagrams.py first) and writes
docs/DengueCAD_Architecture_Slides.pptx. With --pdf, also exports a PDF via
LibreOffice.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DIAGRAMS = ROOT / "docs" / "architecture"
OUT = ROOT / "docs" / "DengueCAD_Architecture_Slides.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TITLE_H = Inches(0.9)
MARGIN = Inches(0.3)
CAPTION_H = Inches(0.5)

NAVY = RGBColor(0x1A, 0x36, 0x5D)
BLUE = RGBColor(0x2C, 0x52, 0x82)
GREY = RGBColor(0x4A, 0x55, 0x68)
LIGHT = RGBColor(0xEB, 0xF4, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _title_bar(slide, title: str, subtitle: str | None = None) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, TITLE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = MARGIN
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = WHITE
    if subtitle:
        run2 = p.add_run()
        run2.text = f"   {subtitle}"
        run2.font.size = Pt(16)
        run2.font.color.rgb = RGBColor(0xBE, 0xE3, 0xF8)


def _footer(slide, n: int) -> None:
    box = slide.shapes.add_textbox(
        MARGIN, SLIDE_H - Inches(0.4), SLIDE_W - 2 * MARGIN, Inches(0.3)
    )
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"DengueCAD architecture  ·  {n}"
    run.font.size = Pt(10)
    run.font.color.rgb = GREY


def _bullets(slide, items: list[str | tuple[str, int]], left, top, width, height, size=24):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = ("•  " if level == 0 else "      –  ") + text
        run.font.size = Pt(size if level == 0 else size - 3)
        run.font.color.rgb = NAVY if level == 0 else GREY
    return box


def _caption(slide, text: str) -> None:
    box = slide.shapes.add_textbox(
        MARGIN, SLIDE_H - Inches(0.35) - CAPTION_H, SLIDE_W - 2 * MARGIN, CAPTION_H
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(14)
    run.font.italic = True
    run.font.color.rgb = GREY


def _picture_fit(slide, path: Path, left, top, width, height) -> None:
    """Insert an image scaled to fit the box, preserving aspect ratio, centred."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(width / iw, height / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(
        str(path), Emu(left + (width - w) // 2), Emu(top + (height - h) // 2), Emu(w), Emu(h)
    )


def _table(slide, rows: list[list[str]], left, top, width, col_widths: list[float], size=14):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.6) * n_rows)
    table = shape.table
    total = sum(col_widths)
    for j, w in enumerate(col_widths):
        table.columns[j].width = Emu(int(width * w / total))
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = text
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(size)
            para.font.bold = i == 0
            para.font.color.rgb = WHITE if i == 0 else NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if i == 0 else (LIGHT if i % 2 else WHITE)
    return shape


def diagram_slide(prs, n, title, image, caption, notes, subtitle=None):
    s = _blank(prs)
    _title_bar(s, title, subtitle)
    top = TITLE_H + Inches(0.1)
    height = SLIDE_H - top - CAPTION_H - Inches(0.4)
    _picture_fit(s, DIAGRAMS / image, MARGIN, top, SLIDE_W - 2 * MARGIN, height)
    _caption(s, caption)
    _footer(s, n)
    _notes(s, notes)


def build(out: Path) -> Path:
    missing = [p for p in (
        "01-system-overview.png", "02-nmprediction-pipeline.png", "03-ga-wrapper.png",
        "04-module-dependencies.png", "05-dengue-data-protocol.png",
        "06-evaluation-flow.png", "07-imputation-benchmark.png", "08-ci-cd.png",
    ) if not (DIAGRAMS / p).is_file()]
    if missing:
        sys.exit(f"Missing diagrams {missing}; run scripts/render_architecture_diagrams.py first.")

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    n = 0

    # 1 — title
    n += 1
    s = _blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.3), SLIDE_W - Inches(1.6), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, color, bold) in enumerate((
        ("DengueCAD", 54, WHITE, True),
        ("System Architecture", 32, RGBColor(0xBE, 0xE3, 0xF8), False),
        ("NMI imputation + NMPrediction (IEEE TITB 2012, Algorithm 1) "
         "for computer-aided diagnosis", 20, WHITE, False),
        (date.today().strftime("%B %Y"), 16, RGBColor(0xA0, 0xAE, 0xC0), False),
    )):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    _notes(s, "DengueCAD is a disease-agnostic CAD stack. This deck walks through its "
              "components, the NMPrediction algorithm, data protocols, evaluation and CI/CD.")

    # 2 — goals
    n += 1
    s = _blank(prs)
    _title_bar(s, "What DengueCAD does")
    _bullets(s, [
        "Implements NMPrediction (Rao & Kumar, IEEE TITB 2012)",
        ("NMI imputation → GA wrapper feature selection → stratified 10-fold ADT", 1),
        ("Reports Accuracy, AUC, Sensitivity (SE), Specificity (SP)", 1),
        "Compares NM against C4.5, Logistic Regression (LOR) and SVM (RBF)",
        ("Table I performance · Table II Wilcoxon test · Table III influential features", 1),
        "Benchmarks imputation: NMI vs MissForest, MICE, kNN, Mean/Median/Mode",
        "Disease-agnostic: dengue CSVs plus 11 KEEL / Weka ARFF disease datasets",
        "Reproducible: scripted experiments, generated PDF report, CI on every push",
    ], MARGIN + Inches(0.3), TITLE_H + Inches(0.4), SLIDE_W - 2 * MARGIN - Inches(0.6),
        SLIDE_H - TITLE_H - Inches(1.2))
    _footer(s, n)
    _notes(s, "The core library does not assume dengue; any tabular dataset with a "
              "decision column can go through the same pipeline.")

    n += 1
    diagram_slide(
        prs, n, "System overview", "01-system-overview.png",
        "Data sources → experiment scripts → denguecad package → CSV results and PDF report; "
        "imputation delegates to the sibling NMI library.",
        "Left to right: the data folder, the scripts that orchestrate each experiment, the "
        "library layers, and outputs. NMI is a separate repository located via NMI_ROOT or ../NMI.",
    )

    # 4 — layers table
    n += 1
    s = _blank(prs)
    _title_bar(s, "Package components", "denguecad/")
    _table(s, [
        ["Layer", "Modules", "Responsibility"],
        ["Data", "data_splits · arff_io · imputation_benchmark", "Load, encode, split (70k / 10k / 20k), ARFF with KEEL ranges"],
        ["Imputation", "nmi_support · imputation_benchmark · missforest_impute", "NMI, MICE, kNN, Mean/Mode, MissForest"],
        ["Algorithm", "nm_prediction · feature_selection", "NMPrediction (Algorithm 1), GA wrapper search"],
        ["Learners", "adt · baselines · classifiers", "ADT, C4.5, LOR, SVM (RBF), classifier factory"],
        ["Evaluation", "metrics · comparison", "Accuracy/AUC/SE/SP, Tables I–III, Wilcoxon"],
        ["Reporting", "scripts/build_comparison_pdf.py", "PDF with legend; red = best, green = NM when not best"],
    ], MARGIN, TITLE_H + Inches(0.5), SLIDE_W - 2 * MARGIN, [1.3, 3.6, 4.6], size=18)
    _footer(s, n)
    _notes(s, "Each layer only depends on layers below it; the evaluation layer is what the "
              "experiment scripts call.")

    n += 1
    diagram_slide(
        prs, n, "NMPrediction pipeline", "02-nmprediction-pipeline.png",
        "Algorithm 1: impute missing values with NMI, pick influential features with a GA "
        "wrapper, then score ADT with stratified 10-fold CV.",
        "Steps 1–2 impute only if values are missing. Step 3 is the GA wrapper. Steps 4–8 pool "
        "fold predictions to compute AUC, SE and SP. The selected features are also used for "
        "the hold-out validation and test sets.",
        subtitle="Algorithm 1",
    )

    n += 1
    diagram_slide(
        prs, n, "GA wrapper feature selection", "03-ga-wrapper.png",
        "Kohavi & John wrapper: the classifier scores each subset. Pc = 1.0, Pm = 0.001, "
        "elitism 2, tournament 3, small subset-size penalty.",
        "Chromosomes are bit masks over features. Fitness is cross-validated accuracy minus a "
        "penalty proportional to subset size, which favours compact subsets.",
        subtitle="Section C",
    )

    n += 1
    diagram_slide(
        prs, n, "Module dependencies", "04-module-dependencies.png",
        "Arrows point to imported modules. Only nmi_support touches the external NMI library.",
        "comparison orchestrates NM and baselines. nm_prediction depends on feature_selection, "
        "baselines, metrics and nmi_support. adt is shared by baselines and the GA classifier factory.",
    )

    n += 1
    diagram_slide(
        prs, n, "Data protocol: dengue.csv", "05-dengue-data-protocol.png",
        "70,000-row single training set · 10,000 complete-case validation · "
        "20,000 complete-case test.",
        "The first 70k rows go through NMI and NMPrediction as one dataset. From the rest, "
        "records with any missing value are dropped; the first 10k complete records are "
        "validation and the next 20k are the test set used to report performance.",
    )

    n += 1
    diagram_slide(
        prs, n, "NM vs baselines evaluation", "06-evaluation-flow.png",
        "All methods share the same folds; per-fold accuracies feed the Wilcoxon signed-rank test.",
        "Table I reports pooled metrics, Table II the Wilcoxon test of NM against each "
        "baseline, Table III the features NM selected. Everything is collected in the "
        "comparison PDF.",
    )

    n += 1
    diagram_slide(
        prs, n, "Imputation benchmark", "07-imputation-benchmark.png",
        "Mask a complete dataset with MCAR missingness, impute with each method, "
        "score against ground truth and downstream classifiers.",
        "Used for dataset.csv and the canonical MissForest Iris demo (20% MCAR). Categorical "
        "cells must match exactly; real values within 5% relative tolerance.",
    )

    # 11 — runners table
    n += 1
    s = _blank(prs)
    _title_bar(s, "Experiment runners and outputs", "scripts/ → results/")
    _table(s, [
        ["Script", "Dataset", "Output"],
        ["run_nm_comparison.py", "dengue.csv (70k / 10k / 20k)", "Tables I–III, hold-out metrics, split summary"],
        ["run_nm_on_dataset_csv.py", "dataset.csv (Bangladesh)", "results/nmprediction_dataset_csv/"],
        ["run_nm_on_clinical.py", "Dengue_clinical_dataset.csv", "results/nmprediction_clinical/"],
        ["run_arff_disease_benchmark.py", "11 disease ARFFs", "results/arff_disease_comparison/"],
        ["run_imputation_comparison.py", "dataset.csv + MCAR", "results/imputation_dataset_csv/"],
        ["run_missforest_comparison.py", "Iris + 20% MCAR", "results/missforest_comparison/"],
        ["build_comparison_pdf.py", "all results", "Performance comparison PDF (docs/)"],
    ], MARGIN, TITLE_H + Inches(0.5), SLIDE_W - 2 * MARGIN, [3.2, 3.2, 4.2], size=17)
    _footer(s, n)
    _notes(s, "results/ is git-ignored; the PDF report is the committed record of performance.")

    n += 1
    diagram_slide(
        prs, n, "CI/CD", "08-ci-cd.png",
        "GitHub Actions on every push and pull request; a green push to main cuts a semver release.",
        "The same gates run locally with scripts/ci_local.sh before pushing. System tests "
        "clone the NMI library and use the datasets in data/.",
        subtitle="GitHub Actions",
    )

    # 13 — design decisions
    n += 1
    s = _blank(prs)
    _title_bar(s, "Design decisions and extension points")
    _bullets(s, [
        "NMI kept as an external library behind nmi_support",
        ("Swap the checkout with NMI_ROOT; no code changes", 1),
        "Pluggable learners through factories (baselines, classifiers)",
        ("GA wrapper can evaluate with ADT, SVM or decision tree", 1),
        "Disease-agnostic data layer",
        ("New datasets: add a CSV/ARFF loader, reuse run_dataset_comparison", 1),
        "Reproducibility",
        ("Fixed random seeds, scripted runs, PDF and diagrams regenerated from source", 1),
    ], MARGIN + Inches(0.3), TITLE_H + Inches(0.4), SLIDE_W - 2 * MARGIN - Inches(0.6),
        SLIDE_H - TITLE_H - Inches(1.2))
    _footer(s, n)
    _notes(s, "These are the seams to use when adding a disease dataset, a new imputer or a new "
              "classifier.")

    # 14 — how to run
    n += 1
    s = _blank(prs)
    _title_bar(s, "Running and regenerating")
    box = s.shapes.add_textbox(
        MARGIN + Inches(0.3), TITLE_H + Inches(0.4), SLIDE_W - 2 * MARGIN - Inches(0.6), Inches(5)
    )
    tf = box.text_frame
    tf.word_wrap = True
    lines = [
        "source .venv/bin/activate",
        "python scripts/run_nm_comparison.py            # dengue.csv experiment",
        "python scripts/run_arff_disease_benchmark.py   # multi-disease ARFFs",
        "python scripts/build_comparison_pdf.py         # performance report",
        "python scripts/render_architecture_diagrams.py # diagrams from ARCHITECTURE.md",
        "python scripts/build_architecture_slides.py    # this deck",
        "bash scripts/ci_local.sh                       # CI gates before push",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = line
        run.font.name = "Courier New"
        run.font.size = Pt(18)
        run.font.color.rgb = NAVY
    _footer(s, n)
    _notes(s, "Diagrams live as Mermaid in docs/ARCHITECTURE.md; images and slides are generated.")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def export_pdf(pptx: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice not found; cannot export PDF.")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
        check=True,
        stdout=subprocess.DEVNULL,
    )
    return pptx.with_suffix(".pdf")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the architecture slide deck")
    parser.add_argument("--out", default=str(OUT))
    parser.add_argument("--pdf", action="store_true", help="Also export a PDF via LibreOffice")
    args = parser.parse_args()
    out = build(Path(args.out))
    print(f"Wrote {out}")
    if args.pdf:
        print(f"Wrote {export_pdf(out)}")


if __name__ == "__main__":
    main()
